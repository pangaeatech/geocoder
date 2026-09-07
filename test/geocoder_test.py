#!/usr/bin/python3
# -.- coding: utf-8 -.-
# -.- dependencies: Python 3.8+ -.-

"""
Geocoder Tests

Copyright (c) 2026 Pangaea Information Technologies, Ltd.
"""

import json

import openpyxl
import pytest

from src import api
from src import census
from src import geocoder
from src import postprocess
from src import preprocess
from src.api import GeocodeResult, Provider, SourceRecord
from src.geocoder import Options, detect_columns, main, process_workbook, write_output_sheet


class MockProvider(Provider):
    """In-test provider: echoes each record so read->geocode->write is exercised."""

    name = "mock"
    requires_key = False

    def geocode(self, records):
        """Returns a fixed result per record, echoing the source address fields."""
        results = []
        for record in records:
            results.append(
                GeocodeResult(
                    result_address=record.address,
                    result_city=record.city,
                    result_stateprov=record.stateprov,
                    result_country=record.country,
                    latitude="40.0",
                    longitude="-75.0",
                    match_type="exact",
                    accuracy=100,
                    location_type="rooftop",
                    match_notes="",
                    raw={"key": record.internal_key, "q": record.address_string()},
                )
            )
        return results


def _make_workbook(path, sheets):
    """
    Writes a temporary workbook for a test.

    Parameters
    ----------
    path
        The path of the workbook to create.
    sheets
        A dict of sheet_name -> list of rows, where the first row is the header.
    """
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        worksheet = workbook.create_sheet(title=name)
        for row in rows:
            worksheet.append(row)
    workbook.save(path)


def test_detect_columns_synonyms_case_insensitive():
    """Header synonyms map to canonical fields regardless of case."""
    header = [
        "Facility ID",
        "NAME",
        "Street Address",
        "city",
        "State/Province",
        "Zip Code",
        "Country",
        "Lat",
        "Lng",
    ]
    mapping = detect_columns(header)
    assert mapping == {
        "ID": 0,
        "NAME": 1,
        "ADDRESS": 2,
        "CITY": 3,
        "STATEPROV": 4,
        "POSTALCODE": 5,
        "COUNTRY": 6,
        "LATITUDE": 7,
        "LONGITUDE": 8,
    }


def test_detect_columns_first_match_wins():
    """When several columns match one field, the first column wins."""
    mapping = detect_columns(["address", "street address"])
    assert mapping["ADDRESS"] == 0


def test_process_workbook_end_to_end(tmp_path):
    """Read, geocode, and write produces the full column schema and values."""
    infile = tmp_path / "in.xlsx"
    outfile = tmp_path / "out.xlsx"
    _make_workbook(
        infile,
        {
            "Sheet1": [
                ["Address", "City", "State", "Zip Code"],
                ["1600 Pennsylvania Ave NW", "Washington", "DC", "20500"],
            ],
        },
    )

    process_workbook(str(infile), str(outfile), MockProvider())

    result = openpyxl.load_workbook(outfile)
    assert result.sheetnames == ["Sheet1"]
    sheet = result["Sheet1"]
    header = [cell.value for cell in sheet[1]]
    assert header == (geocoder.SOURCE_HEADERS + geocoder.RESULT_HEADERS + geocoder.META_HEADERS)
    assert geocoder.DEBUG_HEADER not in header
    row = [cell.value for cell in sheet[2]]
    values = dict(zip(header, row))
    assert values["SOURCE_ADDRESS"] == "1600 Pennsylvania Ave NW"
    assert values["SOURCE_CITY"] == "Washington"
    assert values["RESULT_LATITUDE"] == "40.0"
    assert values["GEOCODER_API"] == "mock"
    assert values["MATCH_TYPE"] == "exact"
    assert values["ACCURACY"] == 100


def test_debug_adds_raw_match_column(tmp_path):
    """The --debug flag appends a RAW_MATCH column with the raw provider JSON."""
    infile = tmp_path / "in.xlsx"
    outfile = tmp_path / "out.xlsx"
    _make_workbook(
        infile,
        {
            "Sheet1": [
                ["Address", "City", "State"],
                ["1 Main St", "Anytown", "CA"],
            ],
        },
    )

    process_workbook(str(infile), str(outfile), MockProvider(), options=Options(debug=True))

    sheet = openpyxl.load_workbook(outfile)["Sheet1"]
    header = [cell.value for cell in sheet[1]]
    assert header[-1] == geocoder.DEBUG_HEADER
    raw_value = sheet[2][len(header) - 1].value
    assert json.loads(raw_value)["key"] == 0


def test_missing_required_columns_sheet_skipped(tmp_path):
    """Sheets lacking a required column are skipped, others are still written."""
    infile = tmp_path / "in.xlsx"
    outfile = tmp_path / "out.xlsx"
    _make_workbook(
        infile,
        {
            "Good": [
                ["Address", "City", "State"],
                ["1 Main St", "Anytown", "CA"],
            ],
            "Bad": [
                ["Name", "City", "State"],
                ["No address column", "Anytown", "CA"],
            ],
        },
    )

    process_workbook(str(infile), str(outfile), MockProvider())

    result = openpyxl.load_workbook(outfile)
    assert result.sheetnames == ["Good"]


def test_all_sheets_skipped_raises(tmp_path):
    """When no sheet qualifies, nothing is written and SystemExit is raised."""
    infile = tmp_path / "in.xlsx"
    outfile = tmp_path / "out.xlsx"
    _make_workbook(
        infile,
        {
            "Bad": [
                ["Name", "City", "State"],
                ["x", "y", "z"],
            ],
        },
    )

    with pytest.raises(SystemExit):
        process_workbook(str(infile), str(outfile), MockProvider())
    assert not outfile.exists()


def test_single_worksheet_selection(tmp_path):
    """Naming a worksheet processes only that sheet."""
    infile = tmp_path / "in.xlsx"
    outfile = tmp_path / "out.xlsx"
    _make_workbook(
        infile,
        {
            "One": [["Address", "City", "State"], ["1 A St", "Town", "CA"]],
            "Two": [["Address", "City", "State"], ["2 B St", "City", "NY"]],
        },
    )

    process_workbook(str(infile), str(outfile), MockProvider(), worksheet="Two")

    result = openpyxl.load_workbook(outfile)
    assert result.sheetnames == ["Two"]


def test_unknown_worksheet_raises(tmp_path):
    """Naming a worksheet that does not exist raises SystemExit."""
    infile = tmp_path / "in.xlsx"
    outfile = tmp_path / "out.xlsx"
    _make_workbook(infile, {"One": [["Address", "City", "State"], ["1 A St", "T", "CA"]]})

    with pytest.raises(SystemExit):
        process_workbook(str(infile), str(outfile), MockProvider(), worksheet="Nope")


def test_country_per_sheet_fills_blank(tmp_path):
    """--countryPerSheet fills blank country cells with the sheet name only."""
    infile = tmp_path / "in.xlsx"
    outfile = tmp_path / "out.xlsx"
    _make_workbook(
        infile,
        {
            "Canada": [
                ["Address", "City", "State", "Country"],
                ["1 Blank St", "Toronto", "ON", ""],
                ["2 Set St", "Paris", "IDF", "France"],
            ],
        },
    )

    process_workbook(str(infile), str(outfile), MockProvider(), options=Options(country_per_sheet=True))

    sheet = openpyxl.load_workbook(outfile)["Canada"]
    header = [cell.value for cell in sheet[1]]
    country_col = header.index("SOURCE_COUNTRY")
    assert sheet[2][country_col].value == "Canada"
    assert sheet[3][country_col].value == "France"


def test_main_missing_infile(tmp_path):
    """A missing input file exits before any provider work."""
    with pytest.raises(SystemExit):
        main([str(tmp_path / "nope.xlsx"), str(tmp_path / "out.xlsx"), "--api", "mock"])


def test_main_outfile_exists(tmp_path):
    """An existing output file exits without overwriting it."""
    infile = tmp_path / "in.xlsx"
    outfile = tmp_path / "out.xlsx"
    _make_workbook(infile, {"S": [["Address", "City", "State"], ["1 A", "T", "CA"]]})
    outfile.write_text("existing")

    with pytest.raises(SystemExit):
        main([str(infile), str(outfile), "--api", "mock"])


def test_main_unknown_api(tmp_path):
    """An api with no registered provider exits cleanly."""
    infile = tmp_path / "in.xlsx"
    _make_workbook(infile, {"S": [["Address", "City", "State"], ["1 A", "T", "CA"]]})

    with pytest.raises(SystemExit):
        main([str(infile), str(tmp_path / "out.xlsx"), "--api", "nonexistent"])


def test_main_runs_registered_provider(tmp_path):
    """A registered provider runs end to end and writes the output file."""
    infile = tmp_path / "in.xlsx"
    outfile = tmp_path / "out.xlsx"
    _make_workbook(infile, {"S": [["Address", "City", "State"], ["1 A St", "Town", "CA"]]})

    api.PROVIDERS["mock"] = MockProvider
    try:
        main([str(infile), str(outfile), "--api", "mock"])
    finally:
        del api.PROVIDERS["mock"]

    assert outfile.exists()


def test_preprocess_adds_flag_columns(tmp_path):
    """--preProcess inserts PRE_ columns between the source and result columns."""
    infile = tmp_path / "in.xlsx"
    outfile = tmp_path / "out.xlsx"
    _make_workbook(
        infile,
        {
            "Sheet1": [
                ["Address", "City", "State"],
                ["PO Box 12", "Anytown", "CA"],
            ],
        },
    )

    process_workbook(str(infile), str(outfile), MockProvider(), options=Options(preprocess=True))

    sheet = openpyxl.load_workbook(outfile)["Sheet1"]
    header = [cell.value for cell in sheet[1]]
    assert header == (geocoder.SOURCE_HEADERS + preprocess.PRE_HEADERS + geocoder.RESULT_HEADERS + geocoder.META_HEADERS)
    values = dict(zip(header, [cell.value for cell in sheet[2]]))
    assert "PO_BOX" in values["PRE_FLAGS"]
    assert "PO Box" in values["PRE_NOTES"]


def test_compare_adds_match_columns(tmp_path):
    """--compare appends a graded MATCH_ column for each compared field."""
    infile = tmp_path / "in.xlsx"
    outfile = tmp_path / "out.xlsx"
    _make_workbook(
        infile,
        {
            "Sheet1": [
                ["Address", "City", "State", "Lat", "Lng"],
                ["1 Main St", "Anytown", "CA", "40.0", "-75.0"],
            ],
        },
    )

    process_workbook(str(infile), str(outfile), MockProvider(), options=Options(compare=True))

    sheet = openpyxl.load_workbook(outfile)["Sheet1"]
    header = [cell.value for cell in sheet[1]]
    assert header[-len(postprocess.MATCH_HEADERS) :] == postprocess.MATCH_HEADERS
    values = dict(zip(header, [cell.value for cell in sheet[2]]))
    assert values["MATCH_ADDRESS"] == "EXACT"
    assert values["MATCH_COUNTRY"] == "BLANK"
    assert values["MATCH_DISTANCE_M"] == 0


def test_api_none_writes_source_and_flags_only(tmp_path):
    """--api none stops after pre-processing and writes no result columns."""
    infile = tmp_path / "in.xlsx"
    outfile = tmp_path / "out.xlsx"
    _make_workbook(infile, {"S": [["Address", "City", "State"], ["PO Box 12", "Anytown", "CA"]]})

    main([str(infile), str(outfile), "--api", "none"])

    sheet = openpyxl.load_workbook(outfile)["S"]
    header = [cell.value for cell in sheet[1]]
    assert header == geocoder.SOURCE_HEADERS + preprocess.PRE_HEADERS
    assert "PO_BOX" in sheet[2][len(header) - 2].value


def test_compare_with_api_none_rejected(tmp_path):
    """--compare has nothing to compare against without a provider, so it exits."""
    infile = tmp_path / "in.xlsx"
    _make_workbook(infile, {"S": [["Address", "City", "State"], ["1 A St", "Town", "CA"]]})

    with pytest.raises(SystemExit):
        main([str(infile), str(tmp_path / "out.xlsx"), "--api", "none", "--compare"])


def test_write_output_sheet_length_mismatch_raises():
    """A provider returning the wrong number of results is rejected."""
    records = [SourceRecord(internal_key=0, address="1 A St")]
    with pytest.raises(ValueError):
        write_output_sheet(openpyxl.Workbook(), "S", records, [], [], "mock", Options())


def test_write_output_sheet_flag_mismatch_raises():
    """Pre-check flags that do not line up with the rows are rejected."""
    records = [SourceRecord(internal_key=0, address="1 A St")]
    with pytest.raises(ValueError):
        write_output_sheet(openpyxl.Workbook(), "S", records, [], None, "none", Options(preprocess=True))


def test_country_per_sheet_reports_no_missing_country(tmp_path, capsys):
    """A country taken from the sheet name is not reported as a missing column."""
    infile = tmp_path / "in.xlsx"
    outfile = tmp_path / "out.xlsx"
    _make_workbook(infile, {"Canada": [["Address", "City", "State"], ["1 A St", "Toronto", "ON"]]})

    process_workbook(str(infile), str(outfile), MockProvider(), options=Options(preprocess=True, country_per_sheet=True))

    assert "COUNTRY" not in capsys.readouterr().out
    sheet = openpyxl.load_workbook(outfile)["Canada"]
    header = [cell.value for cell in sheet[1]]
    assert not dict(zip(header, [cell.value for cell in sheet[2]]))["PRE_FLAGS"]


def test_main_runs_census_provider(tmp_path, monkeypatch):
    """--api census runs end to end and writes the census result columns."""
    infile = tmp_path / "in.xlsx"
    outfile = tmp_path / "out.xlsx"
    _make_workbook(infile, {"S": [["Address", "City", "State"], ["1 Main St", "Town", "CA"]]})

    class _Response:
        """Minimal requests.Response stand-in returning a fixed census row."""

        text = '"0","1 Main St, Town, CA","Match","Exact","1 MAIN ST, TOWN, CA, 90210","-118.0,34.0","1","L","06","037","1","1"\r\n'

        def raise_for_status(self):
            """Mimics a successful response by never raising."""

    def fake_post(*_args, **_kwargs):
        return _Response()

    monkeypatch.setattr(census.requests, "post", fake_post)

    main([str(infile), str(outfile), "--api", "census"])

    sheet = openpyxl.load_workbook(outfile)["S"]
    header = [cell.value for cell in sheet[1]]
    values = dict(zip(header, [cell.value for cell in sheet[2]]))
    assert values["GEOCODER_API"] == "census"
    assert values["RESULT_LATITUDE"] == "34.0"
    assert values["RESULT_LONGITUDE"] == "-118.0"
    assert values["MATCH_TYPE"] == "exact"
