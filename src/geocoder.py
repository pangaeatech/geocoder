#!/usr/bin/python3
# -.- coding: utf-8 -.-
# -.- dependencies: Python 3.8+ -.-

"""
Geocoder

Copyright (c) 2026 Pangaea Information Technologies, Ltd.
"""

import argparse
import json
import os
from dataclasses import dataclass, replace
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook

from .api import (
    KEY_ENV_VARS,
    PROVIDERS,
    AccuracyLevel,
    GeocodeResult,
    Provider,
    SourceRecord,
    resolve_api_key,
)
from .flags import format_flags
from .postprocess import MATCH_HEADERS, compare_record
from .preprocess import BLANK_CHECKED_FIELDS, PRE_HEADERS, QUERY_FIELDS, check_records

CANONICAL_FIELDS = [
    "ID",
    "NAME",
    "ADDRESS",
    "CITY",
    "STATEPROV",
    "POSTALCODE",
    "COUNTRY",
    "LATITUDE",
    "LONGITUDE",
]

COLUMN_SYNONYMS = {
    "ID": ["id", "facility id", "locationid", "location id", "site id", "record id"],
    "NAME": ["name", "facility name", "site name", "location name", "business name"],
    "ADDRESS": ["address", "street address", "address 1", "address1", "addr", "street"],
    "CITY": ["city", "town", "municipality"],
    "STATEPROV": [
        "state",
        "province",
        "state/province",
        "stateprovince",
        "state province",
        "state/prov",
        "prov",
    ],
    "POSTALCODE": ["zipcode", "zip code", "postal code", "postalcode", "zip", "postcode", "postal"],
    "COUNTRY": ["country", "country code"],
    "LATITUDE": ["lat", "latitude"],
    "LONGITUDE": ["lng", "longitude", "long", "lon"],
}

REQUIRED_FIELDS = ["ADDRESS", "CITY", "STATEPROV"]

SOURCE_HEADERS = [f"SOURCE_{field_name}" for field_name in CANONICAL_FIELDS]
RESULT_HEADERS = [f"RESULT_{field_name}" for field_name in CANONICAL_FIELDS]
RESULT_ATTRIBUTES = {
    "ID": "result_id",
    "NAME": "result_name",
    "ADDRESS": "result_address",
    "CITY": "result_city",
    "STATEPROV": "result_stateprov",
    "POSTALCODE": "result_postalcode",
    "COUNTRY": "result_country",
    "LATITUDE": "latitude",
    "LONGITUDE": "longitude",
}
META_ATTRIBUTES = {
    "MATCH_TYPE": "match_type",
    "ACCURACY": "accuracy",
    "LOCATION_TYPE": "location_type",
    "MATCH_NOTES": "match_notes",
}
META_HEADERS = ["GEOCODER_API"] + list(META_ATTRIBUTES)
DEBUG_HEADER = "RAW_MATCH"

NO_API = "none"

NOTHING_TO_WRITE = (
    "error: nothing to write; --compare grades results against their source, and this workbook holds none. "
    f"Geocode it first (--api census --preProcess), then --api {NO_API} --compare the file that produces."
)


@dataclass
class Options:
    """The optional processing and output steps selected on the command line."""

    country_per_sheet: bool = False
    preprocess: bool = False
    compare: bool = False
    debug: bool = False


def load_dotenv(path: str = ".env") -> None:
    """
    Loads KEY=VALUE pairs from a .env file into the environment.

    Existing environment variables are left untouched, and a missing file is ignored.

    Parameters
    ----------
    path : str
        The path of the .env file to read.
    """
    if not os.path.exists(path):
        return

    with open(path, encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, _, value = line.partition("=")
            os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def clean_cell(value) -> str:
    """Converts the specified cell value into a trimmed string."""
    if value is None:
        return ""

    return str(value).strip()


def cell_value(row, index: Optional[int]) -> str:
    """Reads one trimmed cell from a row, or a blank when the column is absent."""
    return clean_cell(row[index]) if index is not None and index < len(row) else ""


def detect_columns(header_cells) -> Dict[str, int]:
    """
    Maps canonical field names to their 0-based column index in the header row.

    Matching is case-insensitive and uses COLUMN_SYNONYMS; when a field has more
    than one matching column, the first one wins.

    Parameters
    ----------
    header_cells
        The header row values to search through.

    Return
    ----------
    mapping : Dict[str, int]
        Each canonical field name mapped to its column index in the row.
    """
    synonym_to_field = {synonym: field_name for field_name, synonyms in COLUMN_SYNONYMS.items() for synonym in synonyms}

    mapping = {}
    for index, value in enumerate(header_cells):
        if value is None:
            continue
        field_name = synonym_to_field.get(str(value).strip().lower())
        if field_name and field_name not in mapping:
            mapping[field_name] = index
    return mapping


def read_sheet_records(rows, column_map: Dict[str, int], sheet_name: str, country_per_sheet: bool) -> List[SourceRecord]:
    """
    Builds SourceRecords from the data rows of a worksheet.

    The header row must already be consumed before calling this. Fully blank rows
    are skipped, and each kept record is assigned a sequential internal key.

    Parameters
    ----------
    rows
        An iterable of data rows (tuples of cell values), with the header excluded.
    column_map : Dict[str, int]
        Each canonical field name mapped to its column index.
    sheet_name : str
        The worksheet name, used as COUNTRY when country_per_sheet is set.
    country_per_sheet : bool
        When True, blank country cells are filled with sheet_name.

    Return
    ----------
    records : List[SourceRecord]
        One record per non-blank data row.
    """
    records = []
    for row in rows:
        values = {field_name: cell_value(row, column_map.get(field_name)) for field_name in CANONICAL_FIELDS}

        if not any(values.values()):
            continue

        if country_per_sheet and not values["COUNTRY"]:
            values["COUNTRY"] = sheet_name

        attributes = {name.lower(): value for name, value in values.items()}
        records.append(SourceRecord(internal_key=len(records), **attributes))
    return records


def blank_fields(column_map: Dict[str, int], sheet_name: str, country_per_sheet: bool = False) -> List[str]:
    """
    Chooses the fields whose blank cells are worth flagging on a worksheet.

    A field the sheet has no column for is blank on every row, which says
    something about the sheet rather than about any one row, so it is reported
    once on stdout and left out of the per-row checks. Only a column the
    provider is given is worth reporting: a sheet with no NAME column geocodes
    exactly as well as one with it, since no query carries a name. COUNTRY
    counts as filled when the sheet name supplies it.

    Parameters
    ----------
    column_map : Dict[str, int]
        Each canonical field name found in the header, mapped to its column.
    sheet_name : str
        The worksheet name, used when reporting the columns it lacks.
    country_per_sheet : bool
        When True, every row takes its country from the sheet name.

    Return
    ----------
    fields : List[str]
        The address-forming fields the sheet provides a value for.
    """
    filled = set(column_map) | ({"COUNTRY"} if country_per_sheet else set())
    present = [name for name in BLANK_CHECKED_FIELDS if name.upper() in filled]
    missing = [name.upper() for name in QUERY_FIELDS if name.upper() not in filled]
    if missing:
        print(f"sheet '{sheet_name}': no {', '.join(missing)} column; every row is missing it")
    return present


def build_headers(options: Options, include_results: bool) -> List[str]:
    """
    Builds the header row for the column sections the run produces.

    Parameters
    ----------
    options : Options
        The processing steps selected on the command line.
    include_results : bool
        Whether a provider ran and result columns are therefore written.

    Return
    ----------
    headers : List[str]
        The header labels, in output order.
    """
    headers = list(SOURCE_HEADERS)
    if options.preprocess:
        headers += PRE_HEADERS
    if include_results:
        headers += RESULT_HEADERS + META_HEADERS
        if options.compare:
            headers += MATCH_HEADERS
        if options.debug:
            headers += [DEBUG_HEADER]
    return headers


def result_cells(result: GeocodeResult, api_name: str) -> List:
    """Renders the result and match-metadata cells of one row."""
    values = [getattr(result, RESULT_ATTRIBUTES[field_name]) for field_name in CANONICAL_FIELDS]
    return values + [api_name] + [getattr(result, attribute) for attribute in META_ATTRIBUTES.values()]


def count_flags(counts: Dict[str, int], flags: List[Dict[str, str]]) -> None:
    """Adds a sheet's flags to a running per-flag row count."""
    for record_flags in flags:
        for name in record_flags:
            counts[name] = counts.get(name, 0) + 1


def report_flags(counts: Dict[str, int]) -> None:
    """
    Prints how many rows carried each pre-check flag, commonest first.

    The point of a pre-processing run is to learn what is wrong with a file, so
    the answer is given on stdout rather than left to be counted by hand in the
    workbook that was just written.

    Parameters
    ----------
    counts : Dict[str, int]
        Each flag raised anywhere in the run, mapped to the rows carrying it.
    """
    if not counts:
        return

    print("pre-check flags:")
    for name, count in sorted(counts.items(), key=lambda item: (-item[1], item[0])):
        print(f"  {count:>7}  {name}")


def finish_sheet(worksheet) -> None:
    """
    Leaves a written sheet ready to triage: header frozen and filters armed.

    The output is read by hand in Excel far more often than by a program, and
    the columns this tool adds are there to be sorted and filtered on, so the
    sheet is handed over with the header pinned and a filter on every column.

    Parameters
    ----------
    worksheet
        The finished worksheet, with its header and every data row appended.
    """
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions


def write_output_sheet(
    workbook: Workbook,
    sheet_name: str,
    records: List[SourceRecord],
    flags: List[Dict[str, str]],
    results: Optional[List[GeocodeResult]],
    api_name: str,
    options: Options,
) -> None:
    """
    Writes the source, pre-check, result, and comparison columns for one worksheet.

    Result, comparison, and debug columns are omitted when no provider ran, so a
    pre-processing-only run writes the source rows alongside their flags.

    Parameters
    ----------
    workbook : Workbook
        The destination workbook a new sheet is added to.
    sheet_name : str
        The name of the sheet to create.
    records : List[SourceRecord]
        The source rows, in output order.
    flags : List[Dict[str, str]]
        The pre-check flags, aligned with records; empty when not requested.
    results : Optional[List[GeocodeResult]]
        The geocoding results aligned with records, or None when no provider ran.
    api_name : str
        The provider name recorded in the GEOCODER_API column.
    options : Options
        The processing steps selected on the command line.

    Raises
    ----------
    ValueError
        If the provider or the pre-checks did not produce exactly one entry per
        input record.
    """
    if results is not None and len(results) != len(records):
        raise ValueError(f"provider '{api_name}' returned {len(results)} results " f"for {len(records)} records")
    if options.preprocess and len(flags) != len(records):
        raise ValueError(f"pre-processing produced {len(flags)} flag sets for {len(records)} records")

    worksheet = workbook.create_sheet(title=sheet_name)
    worksheet.append(build_headers(options, results is not None))

    blanks: List[Optional[GeocodeResult]] = [None] * len(records)
    for index, (record, result) in enumerate(zip(records, results if results is not None else blanks)):
        row = [getattr(record, field_name.lower()) for field_name in CANONICAL_FIELDS]
        if options.preprocess:
            row.append(format_flags(flags[index]))
        if result is not None:
            row += result_cells(result, api_name)
            if options.compare:
                row += compare_record(record, result)
            if options.debug:
                row.append(json.dumps(result.raw, ensure_ascii=False) if result.raw else "")
        worksheet.append(row)

    finish_sheet(worksheet)


def process_workbook(
    infile: str,
    outfile: str,
    provider: Optional[Provider],
    worksheet: Optional[str] = None,
    options: Optional[Options] = None,
) -> None:
    """
    Reads an input workbook, geocodes each qualifying sheet, and writes the output.

    Every worksheet is processed unless a single worksheet is named. Sheets missing
    any required column are skipped with a message written to stdout. A None
    provider halts each sheet after its pre-checks and calls no API at all, so a
    comparison it has no results for is dropped rather than graded against an
    empty result, and a run left with no check to write at all is refused.

    Parameters
    ----------
    infile : str
        The path of the workbook to read.
    outfile : str
        The path of the workbook to create.
    provider : Optional[Provider]
        The geocoding provider used to resolve records, or None to skip geocoding.
    worksheet : Optional[str]
        A single worksheet to process; when None, all sheets are processed.
    options : Optional[Options]
        The processing steps selected on the command line, or None for defaults.

    Raises
    ----------
    SystemExit
        If the named worksheet is missing, no sheet had the required columns, or
        the run was left with nothing to write.
    """
    options = options or Options()
    if options.compare and provider is None:
        if not options.preprocess:
            raise SystemExit(NOTHING_TO_WRITE)
        print("no results to compare against; MATCH_ columns not written")

    source = load_workbook(infile, read_only=True, data_only=True)
    try:
        if worksheet is not None and worksheet not in source.sheetnames:
            raise SystemExit(f"error: worksheet '{worksheet}' not found in {infile}")

        sheet_names = [worksheet] if worksheet else source.sheetnames

        output = Workbook()
        output.remove(output.active)

        counts: Dict[str, int] = {}
        processed = 0
        for name in sheet_names:
            rows = source[name].iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                continue

            column_map = detect_columns(header)
            if not all(field_name in column_map for field_name in REQUIRED_FIELDS):
                print(f"skipping sheet '{name}': missing required columns " f"(need {', '.join(REQUIRED_FIELDS)})")
                continue

            records = read_sheet_records(rows, column_map, name, options.country_per_sheet)
            flags = check_records(records, blank_fields(column_map, name, options.country_per_sheet)) if options.preprocess else []
            results = provider.geocode(records) if provider else None
            api_name = provider.name if provider else NO_API
            write_output_sheet(output, name, records, flags, results, api_name, options)
            count_flags(counts, flags)
            processed += 1
    finally:
        source.close()

    if processed == 0:
        raise SystemExit("error: no worksheets had the required columns; nothing written")

    report_flags(counts)
    output.save(outfile)


def locate_headers(labels: List[str], wanted: Dict[str, str]) -> Dict[str, int]:
    """Maps each wanted field to the column its header occupies, where it has one."""
    return {field_name: labels.index(label) for field_name, label in wanted.items() if label in labels}


def detect_output_columns(header_cells) -> Tuple[Dict[str, int], Dict[str, int], Dict[str, int]]:
    """
    Locates the source, result, and match-metadata columns this tool wrote.

    Parameters
    ----------
    header_cells
        The header row values to search through.

    Return
    ----------
    maps : Tuple[Dict[str, int], Dict[str, int], Dict[str, int]]
        The canonical fields found under each result prefix and the metadata
        fields found beside them, each mapped to their columns.
    """
    labels = [clean_cell(value).upper() for value in header_cells]
    source = locate_headers(labels, {name: f"SOURCE_{name}" for name in CANONICAL_FIELDS})
    result = locate_headers(labels, {name: f"RESULT_{name}" for name in CANONICAL_FIELDS})
    return source, result, locate_headers(labels, {name: name for name in META_ATTRIBUTES})


def read_output_row(
    row, index: int, source_map: Dict[str, int], result_map: Dict[str, int], meta_map: Dict[str, int]
) -> Tuple[SourceRecord, GeocodeResult]:
    """
    Rebuilds the source record and geocoded result a written row was made from.

    The match metadata is read back alongside the result fields, because the
    comparison judges a result partly on what the provider said about it: a
    rebuilt result left at its defaults would be graded as freshly accurate
    however coarsely the provider had resolved it. Accuracy is the one metadata
    cell held as a number, and a sheet that lacks the column, or carries a
    blank, leaves it at the score no match scores.

    Parameters
    ----------
    row
        One data row of a workbook this tool wrote.
    index : int
        The row's position, used as the record's internal key.
    source_map : Dict[str, int]
        Each canonical field mapped to its SOURCE_ column.
    result_map : Dict[str, int]
        Each canonical field mapped to its RESULT_ column.
    meta_map : Dict[str, int]
        Each match-metadata field mapped to its column.

    Return
    ----------
    pair : Tuple[SourceRecord, GeocodeResult]
        The two sides of the row, ready to be compared.
    """
    record = SourceRecord(internal_key=index, **{name.lower(): cell_value(row, source_map.get(name)) for name in CANONICAL_FIELDS})

    values = {RESULT_ATTRIBUTES[name]: cell_value(row, column) for name, column in result_map.items()}
    values.update({META_ATTRIBUTES[name]: cell_value(row, column) for name, column in meta_map.items()})
    accuracy = values.get("accuracy", "")
    values["accuracy"] = int(accuracy) if accuracy.isdigit() else AccuracyLevel.NONE
    return record, GeocodeResult(**values)


def extend_headers(header: List[str], names: List[str]) -> Dict[str, int]:
    """
    Finds the column of each named header, appending the ones the sheet lacks.

    Parameters
    ----------
    header : List[str]
        The header row, extended in place with any missing name.
    names : List[str]
        The headers the run needs a column for.

    Return
    ----------
    columns : Dict[str, int]
        Each name mapped to the column it occupies.
    """
    for name in names:
        if name not in header:
            header.append(name)
    return {name: header.index(name) for name in names}


def place_cells(row: List, columns: Dict[str, int], values: List) -> None:
    """Writes values into their columns, padding the row out to reach them."""
    for column, value in zip(columns.values(), values):
        row.extend([""] * (column + 1 - len(row)))
        row[column] = value


def checked_headers(options: Options) -> List[str]:
    """Names the check columns a re-run refreshes, in output order."""
    return (PRE_HEADERS if options.preprocess else []) + (MATCH_HEADERS if options.compare else [])


def filled_columns(columns: Dict[str, int], rows: List[List]) -> Dict[str, int]:
    """
    Keeps the columns some row carries a value in.

    Output this tool wrote has a column for every canonical field, so one empty
    down its whole length marks a field the source never had. Dropping it spares
    every row a blank-cell flag, as a missing column is already spared.

    Parameters
    ----------
    columns : Dict[str, int]
        Each canonical field mapped to the column it occupies.
    rows : List[List]
        The sheet's data rows.

    Return
    ----------
    columns : Dict[str, int]
        The mapping, less the fields no row has a value for.
    """
    return {name: column for name, column in columns.items() if any(cell_value(row, column) for row in rows)}


def check_cells(record: SourceRecord, result: GeocodeResult, flags: Dict[str, str], options: Options) -> List:
    """Renders the pre-check and comparison cells requested for one row."""
    cells = [format_flags(flags)] if options.preprocess else []
    return cells + (compare_record(record, result) if options.compare else [])


def recheck_sheet(rows, sheet_name: str, options: Options) -> Optional[Tuple[List[str], List[List], List[Dict[str, str]]]]:
    """
    Rebuilds one worksheet of written output with its check columns refreshed.

    Every original cell is kept, so the provider columns, the raw match, and any
    earlier flags survive the round trip. A sheet holding no results keeps its
    pre-checks; only the comparison is dropped, and a sheet thereby left with no
    check to write is passed over.

    Parameters
    ----------
    rows
        An iterable of the worksheet's rows, header first.
    sheet_name : str
        The worksheet name, used when reporting what a sheet lacks.
    options : Options
        The checks to re-run.

    Return
    ----------
    sheet : Optional[Tuple[List[str], List[List], List[Dict[str, str]]]]
        The header, the data rows to write, and the pre-check flags behind
        them, or None when the sheet holds no output of this tool to check.
    """
    header = [clean_cell(value) for value in next(rows, ())]
    source_map, result_map, meta_map = detect_output_columns(header)
    if not source_map:
        print(f"skipping sheet '{sheet_name}': no SOURCE_ columns to check")
        return None

    if options.compare and not result_map:
        print(f"sheet '{sheet_name}': no RESULT_ columns; MATCH_ columns not written")
        options = replace(options, compare=False)

    if not options.preprocess and not options.compare:
        return None

    data = [list(row) for row in rows if any(clean_cell(value) for value in row)]
    pairs = [read_output_row(row, index, source_map, result_map, meta_map) for index, row in enumerate(data)]
    flags = check_records([record for record, _ in pairs], blank_fields(filled_columns(source_map, data), sheet_name)) if options.preprocess else []
    columns = extend_headers(header, checked_headers(options))

    for index, (row, pair) in enumerate(zip(data, pairs)):
        place_cells(row, columns, check_cells(*pair, flags[index] if flags else {}, options))
    return header, data, flags


def recheck_workbook(infile: str, outfile: str, worksheet: Optional[str] = None, options: Optional[Options] = None) -> None:
    """
    Re-runs the checks over a workbook this tool already wrote, calling no API.

    The requested check columns are refreshed in place when the sheet already
    has them and appended when it does not, so a file geocoded without them can
    be compared after the fact without paying for the results again.

    Parameters
    ----------
    infile : str
        The path of the workbook to read, as written by an earlier run.
    outfile : str
        The path of the workbook to create.
    worksheet : Optional[str]
        A single worksheet to process; when None, all sheets are processed.
    options : Optional[Options]
        The checks to re-run, or None for defaults.

    Raises
    ----------
    SystemExit
        If the named worksheet is missing, or no sheet held the expected columns.
    """
    options = options or Options()
    source = load_workbook(infile, read_only=True, data_only=True)
    try:
        if worksheet is not None and worksheet not in source.sheetnames:
            raise SystemExit(f"error: worksheet '{worksheet}' not found in {infile}")

        output = Workbook()
        output.remove(output.active)

        counts: Dict[str, int] = {}
        processed = 0
        for name in [worksheet] if worksheet else source.sheetnames:
            sheet = recheck_sheet(source[name].iter_rows(values_only=True), name, options)
            if sheet is None:
                continue

            header, data, flags = sheet
            worksheet_out = output.create_sheet(title=name)
            worksheet_out.append(header)
            for row in data:
                worksheet_out.append(row)
            finish_sheet(worksheet_out)
            count_flags(counts, flags)
            processed += 1
    finally:
        source.close()

    if processed == 0:
        raise SystemExit(NOTHING_TO_WRITE)

    report_flags(counts)
    output.save(outfile)


def holds_output(infile: str, worksheet: Optional[str]) -> bool:
    """
    Reports whether a workbook is one this tool wrote earlier.

    Whether the checks are being added to source rows or refreshed on a file
    already geocoded is a property of the file rather than of the options asked
    for, so it is read off the headers. A worksheet the file does not have is
    left for the run itself to report as missing.

    Parameters
    ----------
    infile : str
        The path of the workbook to read.
    worksheet : Optional[str]
        A single worksheet to look at; when None, every sheet is considered.

    Return
    ----------
    bool
        True when any sheet considered carries this tool's SOURCE_ columns.
    """
    source = load_workbook(infile, read_only=True, data_only=True)
    try:
        names = [worksheet] if worksheet in source.sheetnames else source.sheetnames
        return any(detect_output_columns(next(source[name].iter_rows(values_only=True), ()))[0] for name in names)
    finally:
        source.close()


def parse_args(argv: Optional[List[str]] = None):
    """Takes and parses the command-line arguments given by the user."""
    parser = argparse.ArgumentParser(description="Geocode address rows in an Excel workbook.")

    parser.add_argument("infile", help="Excel file to read (must exist)")
    parser.add_argument("outfile", help="Excel file to create (must not exist)")
    parser.add_argument(
        "worksheet",
        nargs="?",
        default=None,
        help="optional single worksheet to process",
    )
    parser.add_argument(
        "--api",
        default="census",
        metavar="{census|geocodio|google|none}",
        help=f"geocoding provider, or '{NO_API}' to check a workbook without geocoding it (default: census)",
    )
    parser.add_argument(
        "--apiKey",
        dest="api_key",
        default=None,
        help="API key; falls back to .env / environment",
    )
    parser.add_argument(
        "--countryPerSheet",
        dest="country_per_sheet",
        action="store_true",
        help="use the worksheet name as COUNTRY when a row's country is blank",
    )
    parser.add_argument(
        "--preProcess",
        dest="preprocess",
        action="store_true",
        help="add a PRE_FLAGS column flagging problems in the source rows",
    )
    parser.add_argument(
        "--compare",
        action="store_true",
        help="add MATCH_ columns grading each source cell against the result cell; "
        f"with --api {NO_API} it grades a workbook this tool geocoded earlier",
    )
    parser.add_argument(
        "--debug",
        action="store_true",
        help="add a RAW_MATCH column populated with raw provider JSON",
    )
    return parser.parse_args(argv)


def select_provider(args) -> Optional[Provider]:
    """
    Builds the provider named on the command line, or None when it is 'none'.

    Parameters
    ----------
    args
        The parsed command-line arguments.

    Return
    ----------
    Optional[Provider]
        The provider to geocode with, or None to stop after pre-processing.

    Raises
    ----------
    SystemExit
        If the api is unknown, or a required API key is not configured.
    """
    if args.api == NO_API:
        return None

    provider_cls = PROVIDERS.get(args.api)
    if provider_cls is None:
        available = ", ".join(sorted(PROVIDERS) + [NO_API])
        raise SystemExit(f"error: unknown api '{args.api}' (available: {available})")

    api_key = None
    if provider_cls.requires_key:
        api_key = resolve_api_key(args.api, args.api_key)
        if not api_key:
            raise SystemExit(f"error: api '{args.api}' requires an API key " f"(pass --apiKey or set {KEY_ENV_VARS[args.api]})")

    return provider_cls(api_key)


def main(argv: Optional[List[str]] = None) -> None:
    """
    Runs the geocoder: validates arguments, selects a provider, and writes the output.

    Each option asks for one thing and implies no other, so naming no provider
    leaves the checks as the whole of the work and a run with neither has
    nothing to do.

    Parameters
    ----------
    argv : Optional[List[str]]
        The argument list to parse; defaults to sys.argv when None.

    Raises
    ----------
    SystemExit
        If the input file is missing, the output file already exists, the api is
        unknown, a required API key is not configured, or no provider and no
        check was asked for.
    """
    load_dotenv()
    args = parse_args(argv)

    if not os.path.isfile(args.infile):
        raise SystemExit(f"error: input file '{args.infile}' does not exist")
    if os.path.exists(args.outfile):
        raise SystemExit(f"error: output file '{args.outfile}' already exists")

    provider = select_provider(args)
    if provider is None and not (args.preprocess or args.compare):
        raise SystemExit(f"error: --api {NO_API} has nothing to do; add --preProcess or --compare")

    options = Options(
        country_per_sheet=args.country_per_sheet,
        preprocess=args.preprocess,
        compare=args.compare,
        debug=args.debug,
    )
    if provider is None and holds_output(args.infile, args.worksheet):
        recheck_workbook(args.infile, args.outfile, worksheet=args.worksheet, options=options)
    else:
        process_workbook(args.infile, args.outfile, provider, worksheet=args.worksheet, options=options)
    print(f"wrote {args.outfile}")


if __name__ == "__main__":
    main()
